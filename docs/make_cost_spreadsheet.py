"""Generates docs/cost_estimation.xlsx. Edit the blue input cells in Excel; everything else is a formula."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

wb = Workbook()
a = wb.active
a.title = "Assumptions"
blue, bold = Font(color="0000FF"), Font(bold=True)
head = PatternFill("solid", fgColor="DDEBF7")

a["A1"], a["A1"].font = "Cost assumptions (USD per 1,000,000 tokens)", Font(bold=True, size=13)
a["A2"] = "Blue cells are inputs. Gemini paid-tier prices checked on 20 Oct 2026 (ai.google.dev/gemini-api/docs/pricing). The free tier costs 0."
for c, t in zip("ABCD", ["Role", "Model", "Input price", "Output price"]):
    a[f"{c}4"], a[f"{c}4"].font, a[f"{c}4"].fill = t, bold, head
rows = [("Primary", "gemini-3.6-flash", 0.75, 3.75), ("Fallback", "gemini-3.5-flash-lite", 0.30, 2.50)]
for i, r in enumerate(rows, start=5):
    for c, v in zip("ABCD", r):
        a[f"{c}{i}"] = v
        if c in "BCD":
            a[f"{c}{i}"].font = blue
a["A8"], a["B8"] = "Share of traffic served by fallback", 0.05
a["B8"].font, a["B8"].number_format = blue, "0%"
a["A9"], a["B9"] = "Average input tokens per request", 55
a["A10"], a["B10"] = "Average output tokens per request", 120
a["A11"], a["B11"] = "Days per month", 30
for r in (9, 10, 11):
    a[f"B{r}"].font = blue
for r in (5, 6):
    for c in "CD":
        a[f"{c}{r}"].number_format = "$#,##0.00"
a.column_dimensions["A"].width = 38
for c in "BCD":
    a.column_dimensions[c].width = 16

s = wb.create_sheet("Scenarios")
s["A1"], s["A1"].font = "Monthly cost scenarios", Font(bold=True, size=13)
hdr = ["Scenario", "Requests per day", "Primary cost / request", "Fallback cost / request",
       "Blended cost / request", "Cost / day", "Cost / month"]
for i, h in enumerate(hdr, start=1):
    c = s.cell(row=3, column=i, value=h)
    c.font, c.fill, c.alignment = bold, head, Alignment(wrap_text=True, vertical="center")
for i, (name, rpd) in enumerate([("Pilot", 100), ("Small team", 1000), ("Busy product", 10000), ("Large", 100000)], start=4):
    s[f"A{i}"], s[f"B{i}"] = name, rpd
    s[f"B{i}"].font, s[f"B{i}"].number_format = blue, "#,##0"
    s[f"C{i}"] = "=(Assumptions!$B$9*Assumptions!$C$5+Assumptions!$B$10*Assumptions!$D$5)/1000000"
    s[f"D{i}"] = "=(Assumptions!$B$9*Assumptions!$C$6+Assumptions!$B$10*Assumptions!$D$6)/1000000"
    s[f"E{i}"] = f"=(1-Assumptions!$B$8)*C{i}+Assumptions!$B$8*D{i}"
    s[f"F{i}"] = f"=B{i}*E{i}"
    s[f"G{i}"] = f"=F{i}*Assumptions!$B$11"
    for c in "CDE":
        s[f"{c}{i}"].number_format = "$0.000000"
    for c in "FG":
        s[f"{c}{i}"].number_format = "$#,##0.00"
s.column_dimensions["A"].width = 16
for c in "BCDEFG":
    s.column_dimensions[c].width = 18
s.row_dimensions[3].height = 32

n = wb.create_sheet("Notes")
for i, line in enumerate([
    "How the API estimates cost per request:",
    "cost = (input_tokens x input_price + output_tokens x output_price) / 1,000,000",
    "The app logs the real token usage from each response in logs/usage_log.csv, so you can compare estimates to actuals.",
    "Failed attempts return no usage data, so they are not in the logged cost. Timeouts may still be billed by the provider.",
    "Prices checked on 20 Oct 2026 from Google's pricing page. The gemini-3.6-flash price rises on 1 Jan 2027 to 1.50 input and 7.50 output per million tokens.",
    "Output prices include thinking tokens, but the API's reported output_tokens appeared to leave them out in testing, so real paid-tier costs may be higher than estimated.",
    "Average tokens (55 in, 120 out) come from the observed usage log of the support questions used in testing.",
    "The free tier costs 0 but is rate limited and its content may be used by Google to improve its products.",
], start=1):
    n[f"A{i}"] = line
n.column_dimensions["A"].width = 150
wb.save("docs/cost_estimation.xlsx")
